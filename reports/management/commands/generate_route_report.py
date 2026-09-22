from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

from django.core.management.base import BaseCommand
from django.db.models import Count, Q, Sum
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter

from flights.models import Flight
from reports.models import RouteReport


class Command(BaseCommand):
    help = 'Generate route statistics and export them to Excel.'

    def add_arguments(self, parser):
        parser.add_argument(
            '--output',
            default='route_report.xlsx',
            help='Output Excel file path.'
        )

    def handle(self, *args, **options):
        output_path = Path(options['output'])

        # ---------------------------------------------------------
        # 1. Flight statistics
        # ---------------------------------------------------------
        flight_stats = (
            Flight.objects
            .values(
                'departure',
                'destination',
            )
            .annotate(
                total_flights=Count('id'),
                total_seats=Sum('airplane__capacity'),
            )
            .order_by(
                'departure',
                'destination',
            )
        )

        # ---------------------------------------------------------
        # 2. Active reservation statistics
        # ---------------------------------------------------------
        reservation_stats = (
            Flight.objects
            .values(
                'departure',
                'destination',
            )
            .annotate(
                used_seats=Count(
                    'reservations',
                    filter=Q(
                        reservations__status=True
                    )
                )
            )
            .order_by(
                'departure',
                'destination',
            )
        )

        reservation_map = {
            (
                item['departure'],
                item['destination'],
            ): item['used_seats']
            for item in reservation_stats
        }

        # ---------------------------------------------------------
        # 3. Most used airplane model per route
        # ---------------------------------------------------------
        model_stats = (
            Flight.objects
            .values(
                'departure',
                'destination',
                'airplane__model',
            )
            .annotate(
                model_count=Count('id')
            )
            .order_by(
                'departure',
                'destination',
                '-model_count',
                'airplane__model',
            )
        )

        most_used_model_map = {}

        for item in model_stats:
            route_key = (
                item['departure'],
                item['destination'],
            )

            # First record for each route is the most used model
            # because the queryset is ordered by model_count DESC.
            if route_key not in most_used_model_map:
                most_used_model_map[route_key] = (
                    item['airplane__model']
                )

        # ---------------------------------------------------------
        # 4. Prepare report rows
        # ---------------------------------------------------------
        report_rows = []

        for item in flight_stats:
            route_key = (
                item['departure'],
                item['destination'],
            )

            total_flights = item['total_flights']
            total_seats = item['total_seats'] or 0
            used_seats = reservation_map.get(route_key, 0)

            if total_seats > 0:
                occupancy_rate = (
                    Decimal(used_seats)
                    / Decimal(total_seats)
                    * Decimal('100')
                ).quantize(
                    Decimal('0.01'),
                    rounding=ROUND_HALF_UP
                )
            else:
                occupancy_rate = Decimal('0.00')

            most_used_plane_model = (
                most_used_model_map.get(route_key)
                or 'N/A'
            )

            report_rows.append({
                'departure': item['departure'],
                'destination': item['destination'],
                'total_flights': total_flights,
                'total_seats': total_seats,
                'used_seats': used_seats,
                'occupancy_rate': occupancy_rate,
                'most_used_plane_model': most_used_plane_model,
            })

        # ---------------------------------------------------------
        # 5. Persist / upsert RouteReport records
        # ---------------------------------------------------------
        route_reports = [
            RouteReport(
                departure=row['departure'],
                destination=row['destination'],
                total_flights=row['total_flights'],
                total_seats=row['total_seats'],
                used_seats=row['used_seats'],
                occupancy_rate=row['occupancy_rate'],
                most_used_plane_model=row[
                    'most_used_plane_model'
                ],
            )
            for row in report_rows
        ]

        if route_reports:
            RouteReport.objects.bulk_create(
                route_reports,
                update_conflicts=True,
                update_fields=[
                    'total_flights',
                    'total_seats',
                    'used_seats',
                    'occupancy_rate',
                    'most_used_plane_model',
                ],
                unique_fields=[
                    'departure',
                    'destination',
                ],
            )

        # ---------------------------------------------------------
        # 6. Create Excel workbook
        # ---------------------------------------------------------
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = 'Route Report'

        headers = [
            'Departure',
            'Destination',
            'Total Flights',
            'Total Seats',
            'Used Seats',
            'Occupancy Rate (%)',
            'Most Used Plane Model',
        ]

        worksheet.append(headers)

        # Header formatting
        for cell in worksheet[1]:
            cell.font = Font(bold=True)
            cell.alignment = Alignment(
                horizontal='center'
            )

        # Data rows
        for row in report_rows:
            worksheet.append([
                row['departure'],
                row['destination'],
                row['total_flights'],
                row['total_seats'],
                row['used_seats'],
                row['occupancy_rate'],
                row['most_used_plane_model'],
            ])

        # Occupancy rate formatting
        for row_number in range(2, worksheet.max_row + 1):
            worksheet.cell(
                row=row_number,
                column=6
            ).number_format = '0.00'

        # Freeze header row
        worksheet.freeze_panes = 'A2'

        # Enable filtering
        worksheet.auto_filter.ref = worksheet.dimensions

        # Auto-size columns
        for column_index in range(1, worksheet.max_column + 1):
            column_letter = get_column_letter(column_index)

            max_length = 0

            for cell in worksheet[column_letter]:
                if cell.value is not None:
                    max_length = max(
                        max_length,
                        len(str(cell.value))
                    )

            worksheet.column_dimensions[
                column_letter
            ].width = min(max_length + 2, 40)

        # ---------------------------------------------------------
        # 7. Save Excel file
        # ---------------------------------------------------------
        if output_path.parent != Path('.'):
            output_path.parent.mkdir(
                parents=True,
                exist_ok=True
            )

        workbook.save(output_path)

        # ---------------------------------------------------------
        # 8. Console output
        # ---------------------------------------------------------
        self.stdout.write(
            self.style.SUCCESS(
                f'Route report generated successfully: '
                f'{output_path}'
            )
        )

        self.stdout.write(
            f'Total routes: {len(report_rows)}'
        )